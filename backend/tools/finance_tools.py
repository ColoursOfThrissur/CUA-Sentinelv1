import logging
import sqlite3
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional
import httpx

from db.connections import get_operational_db
from core.alert_manager import AlertManager

logger = logging.getLogger(__name__)


class FinanceTools:
    """
    Tools for personal finance portfolio tracking, market data fetching (stocks, crypto, commodities, forex),
    and triggering automated price shift alerts (e.g. Δ >= 5%).
    """

    def __init__(self, alert_manager: Optional[AlertManager] = None):
        self.http_client = httpx.AsyncClient(timeout=10.0, headers={"User-Agent": "Mozilla/5.0"})
        self.alert_manager = alert_manager or AlertManager()

    async def fetch_ticker_quote(self, symbol: str) -> Dict[str, Any]:
        """
        Fetches live ticker data using public finance endpoints (e.g., Yahoo Finance Chart API).
        Supports tickers like: RELIANCE, INFY, TATAMOTORS, AAPL, NVDA, BTC-USD, GC=F (Gold), INR=X.
        Auto-appends .NS for Indian stocks if necessary.
        """
        clean_symbol = symbol.strip().upper()
        symbols_to_try = [clean_symbol]
        if not any(char in clean_symbol for char in [".", "-", "="]):
            symbols_to_try.append(f"{clean_symbol}.NS")

        for sym in symbols_to_try:
            url = f"https://query1.finance.yahoo.com/v8/finance/chart/{sym}?interval=1d&range=2d"
            try:
                res = await self.http_client.get(url)
                if res.status_code == 200:
                    data = res.json()
                    result = data.get("chart", {}).get("result")
                    if result and len(result) > 0:
                        meta = result[0]["meta"]
                        current_price = meta.get("regularMarketPrice", 0.0)
                        previous_close = meta.get("chartPreviousClose", meta.get("previousClose", current_price))

                        change_24h_pct = 0.0
                        if previous_close and previous_close > 0:
                            change_24h_pct = round(((current_price - previous_close) / previous_close) * 100, 2)

                        asset_name = meta.get("shortName") or meta.get("longName") or clean_symbol
                        instrument_type = meta.get("instrumentType", "UNKNOWN").upper()
                        currency = meta.get("currency", "INR")

                        return {
                            "symbol": clean_symbol,
                            "name": asset_name,
                            "price": current_price,
                            "change_24h_pct": change_24h_pct,
                            "currency": currency,
                            "instrument_type": instrument_type,
                            "timestamp": datetime.now(timezone.utc).isoformat(),
                            "success": True,
                        }
            except Exception as e:
                logger.debug(f"Error fetching ticker {sym}: {e}")

        # Check database for existing price if live quote unavailable (e.g. Mutual Funds)
        conn = get_operational_db()
        cur = conn.execute("SELECT buy_price, current_price FROM user_portfolios WHERE asset_symbol = ?", (clean_symbol,))
        row = cur.fetchone()
        conn.close()

        stored_buy_price = float(row["buy_price"]) if row and row["buy_price"] is not None else 0.0
        stored_cur_price = float(row["current_price"]) if row and row["current_price"] is not None else 0.0

        fallback_prices = {
            "BTC-USD": 7800000.00,
            "ETH-USD": 290000.00,
            "NVDA": 11500.00,
            "AAPL": 18800.00,
            "GC=F": 72000.00,
            "INR=X": 83.90,
            "RELIANCE": 2980.50,
            "INFY": 1850.25,
            "TATAMOTORS": 1020.00,
            "TCS": 4200.00,
            "HDFCBANK": 1650.00,
        }
        price = stored_cur_price or stored_buy_price or fallback_prices.get(clean_symbol, 0.0)

        return {
            "symbol": clean_symbol,
            "name": clean_symbol,
            "price": price,
            "change_24h_pct": 0.0,
            "currency": "INR",
            "instrument_type": "STOCK/MUTUAL_FUND",
            "timestamp": datetime.now(timezone.utc).isoformat(),
            "success": False,
            "is_fallback": True,
        }

    def get_portfolio(self) -> List[Dict[str, Any]]:
        conn = get_operational_db()
        cur = conn.execute("SELECT * FROM user_portfolios ORDER BY updated_at DESC")
        rows = [dict(row) for row in cur.fetchall()]
        conn.close()
        return rows

    def upsert_asset(
        self,
        symbol: str,
        asset_type: str,
        name: str,
        quantity: float,
        buy_price: float,
        alert_threshold_pct: float = 5.0,
        notes: str = "",
    ) -> Dict[str, Any]:
        symbol = symbol.strip().upper()
        conn = get_operational_db()
        conn.execute(
            """
            INSERT INTO user_portfolios 
            (asset_symbol, asset_type, asset_name, quantity, buy_price, alert_threshold_pct, notes, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(asset_symbol) DO UPDATE SET
                asset_type = excluded.asset_type,
                asset_name = excluded.asset_name,
                quantity = excluded.quantity,
                buy_price = excluded.buy_price,
                alert_threshold_pct = excluded.alert_threshold_pct,
                notes = excluded.notes,
                updated_at = excluded.updated_at
            """,
            (
                symbol,
                asset_type.upper(),
                name,
                quantity,
                buy_price,
                alert_threshold_pct,
                notes,
                datetime.now(timezone.utc).isoformat(),
            ),
        )
        conn.commit()
        conn.close()
        return {"symbol": symbol, "status": "SAVED"}

    def delete_asset(self, symbol: str) -> bool:
        conn = get_operational_db()
        conn.execute("DELETE FROM user_portfolios WHERE asset_symbol = ?", (symbol.strip().upper(),))
        conn.commit()
        conn.close()
        return True

    def clear_portfolio(self) -> bool:
        conn = get_operational_db()
        conn.execute("DELETE FROM user_portfolios")
        conn.commit()
        conn.close()
        return True

    async def evaluate_portfolio_alerts(self) -> List[Dict[str, Any]]:
        """
        Scans all portfolio assets, updates market prices, and dispatches alerts if Δ >= alert_threshold_pct (e.g. >= 5%).
        """
        assets = self.get_portfolio()
        alerts_triggered = []

        for asset in assets:
            quote = await self.fetch_ticker_quote(asset["asset_symbol"])
            # Only update current price if live quote returned or valid price exists
            current_price = quote["price"] if quote["price"] > 0 else asset.get("buy_price", 0.0)
            change_24h = quote["change_24h_pct"]
            threshold = asset["alert_threshold_pct"]

            conn = get_operational_db()
            conn.execute(
                "UPDATE user_portfolios SET current_price = ?, price_change_24h = ?, updated_at = ? WHERE asset_symbol = ?",
                (current_price, change_24h, datetime.now(timezone.utc).isoformat(), asset["asset_symbol"]),
            )
            conn.commit()
            conn.close()

            if abs(change_24h) >= threshold and change_24h != 0.0:
                severity = "WARNING" if abs(change_24h) < 10.0 else "CRITICAL"
                direction = "📈 SURGED" if change_24h > 0 else "📉 DROPPED"
                title = f"{asset['asset_symbol']} {direction} by {change_24h:+.2f}%"
                message = f"{asset['asset_name']} is currently trading at ₹{current_price:,.2f} ({change_24h:+.2f}% in 24h). Threshold trigger is {threshold}%."

                fields = {
                    "Current Price": f"₹{current_price:,.2f}",
                    "24h Change": f"{change_24h:+.2f}%",
                    "Holdings": f"{asset['quantity']} units",
                    "Total Value": f"₹{(asset['quantity'] * current_price):,.2f}",
                }

                await self.alert_manager.dispatch_alert(
                    title=title,
                    message=message,
                    severity=severity,
                    category="FINANCE",
                    fields=fields,
                )

                alerts_triggered.append({
                    "symbol": asset["asset_symbol"],
                    "price": current_price,
                    "change_24h": change_24h,
                    "threshold": threshold,
                })

        return alerts_triggered

    async def fetch_portfolio_news(self) -> List[Dict[str, Any]]:
        assets = self.get_portfolio()
        if not assets:
            return []

        news_items = []
        try:
            from tools.web_search import WebSearchTool
            ws = WebSearchTool()
            for asset in assets[:5]:
                sym = asset.get("asset_symbol", "")
                name = asset.get("asset_name", sym)
                query = f"{name} {sym} stock market breaking news earnings report"
                res = ws.search(query, max_results=2)
                if res and isinstance(res, str) and len(res.strip()) > 20:
                    news_items.append({
                        "symbol": sym,
                        "name": name,
                        "news_snippet": res[:400]
                    })
        except Exception as err:
            logger.warning(f"Error fetching portfolio news: {err}")

        return news_items

    def import_excel_holdings(self, bytes_data: bytes) -> int:
        """
        Parses Demat holdings & Mutual Fund statements from Excel (.xlsx/.xls) statement exports
        (Zerodha, Groww, Indmoney, AngelOne, Upstox, ICICI, CAMS, Karvy, etc.).
        """
        import io
        import re
        import openpyxl

        try:
            wb = openpyxl.load_workbook(io.BytesIO(bytes_data), data_only=True)
            best_sheet_rows = []
            for sheetname in wb.sheetnames:
                s = wb[sheetname]
                r_list = list(s.iter_rows(values_only=True))
                valid_rows = [r for r in r_list if r and any(cell is not None and str(cell).strip() != "" for cell in r)]
                if len(valid_rows) > len(best_sheet_rows):
                    best_sheet_rows = valid_rows

            rows = best_sheet_rows
            if not rows or len(rows) < 2:
                return 0

            header_idx = -1
            best_score = 0
            best_cols = {"symbol": -1, "qty": -1, "price": -1, "cur_price": -1, "invested": -1, "val": -1, "name": -1, "isin": -1}

            for idx, r in enumerate(rows[:40]):
                if not r:
                    continue
                str_row = [str(cell or "").lower().strip() for cell in r]

                sym_c, qty_c, price_c, cur_c, inv_c, val_c, name_c, isin_c = -1, -1, -1, -1, -1, -1, -1, -1

                for c_idx, val in enumerate(str_row):
                    if not val:
                        continue
                    if any(k == val or k in val for k in ["instrument", "symbol", "ticker", "trading symbol", "stock", "scrip", "security name", "particulars", "company", "scheme", "fund name", "scheme name"]):
                        if sym_c == -1 and "isin" not in val: sym_c = c_idx
                    if any(k in val for k in ["isin", "folio", "folio no"]):
                        if isin_c == -1: isin_c = c_idx
                    if any(k in val for k in ["qty", "quantity", "units", "shares", "holding", "net qty", "cur qty", "total qty", "avail qty", "available units", "balance units"]):
                        if qty_c == -1: qty_c = c_idx
                    if any(k in val for k in ["buy avg", "avg price", "average price", "buy price", "avg cost", "average cost", "cost price", "buy rate", "purchase price", "buy nav", "avg nav", "average nav", "purchase nav"]):
                        if price_c == -1 and "current" not in val and "market" not in val and "val" not in val: price_c = c_idx
                    if any(k in val for k in ["total cost", "invested amount", "invested", "total investment", "cost value", "purchase value", "investment value", "total invested", "inv amount"]):
                        if inv_c == -1: inv_c = c_idx
                    if any(k in val for k in ["ltp", "last price", "current price", "cur price", "market price", "cur nav", "current nav", "nav", "closing nav"]):
                        if cur_c == -1: cur_c = c_idx
                    if any(k in val for k in ["current value", "cur value", "market value", "present value", "total value", "valuation", "curr value", "cur val"]):
                        if val_c == -1: val_c = c_idx
                    if any(k in val for k in ["company name", "stock name", "security name", "description", "scheme name"]):
                        if name_c == -1: name_c = c_idx

                score = 0
                if sym_c != -1 or isin_c != -1 or name_c != -1: score += 2
                if qty_c != -1: score += 2
                if price_c != -1 or inv_c != -1: score += 2
                if cur_c != -1 or val_c != -1: score += 1

                if score > best_score and score >= 3:
                    best_score = score
                    header_idx = idx
                    best_cols = {
                        "symbol": sym_c if sym_c != -1 else (name_c if name_c != -1 else isin_c),
                        "qty": qty_c,
                        "price": price_c,
                        "cur_price": cur_c,
                        "invested": inv_c,
                        "val": val_c,
                        "name": name_c if name_c != -1 else sym_c,
                        "isin": isin_c
                    }

            if header_idx == -1:
                header_idx = 0
                best_cols = {"symbol": 0, "qty": 1, "price": 2, "cur_price": -1, "invested": -1, "val": -1, "name": -1, "isin": -1}

            logger.info(f"Excel Demat import detected header at row {header_idx} (Score {best_score}): {best_cols}")

            symbol_col = best_cols["symbol"]
            qty_col = best_cols["qty"]
            price_col = best_cols["price"]
            cur_price_col = best_cols["cur_price"]
            invested_col = best_cols["invested"]
            val_col = best_cols["val"]
            name_col = best_cols["name"]

            imported_count = 0
            for r in rows[header_idx + 1:]:
                if not r or symbol_col >= len(r) or r[symbol_col] is None:
                    continue

                raw_sym_val = str(r[symbol_col]).strip()
                if not raw_sym_val:
                    continue

                lower_sym = raw_sym_val.lower()
                if any(ignored in lower_sym for ignored in ["total", "subtotal", "summary", "grand total", "disclaimer", "notes", "account", "balance", "instrument", "symbol", "ticker", "particulars"]):
                    continue

                # 1. Parse ticker symbol / scheme name cleanly
                paren_match = re.search(r'\(([^)]+)\)', raw_sym_val)
                if paren_match and 2 <= len(paren_match.group(1).strip()) <= 15:
                    clean_sym = paren_match.group(1).upper().strip()
                else:
                    clean_sym = raw_sym_val.upper().strip()
                    clean_sym = re.sub(r'[\-:]?(EQ|BE|BL|SM|ST)$', '', clean_sym).strip()

                clean_sym = re.sub(r'[^A-Z0-9.\-]', '', clean_sym)
                if not clean_sym or len(clean_sym) < 2:
                    continue

                # 2. Parse Quantity
                raw_qty = 1.0
                if qty_col != -1 and qty_col < len(r) and r[qty_col] is not None:
                    try:
                        q_str = re.sub(r'[^0-9.]', '', str(r[qty_col]))
                        if q_str:
                            raw_qty = float(q_str)
                    except ValueError:
                        raw_qty = 1.0

                # 3. Parse Buy Price / NAV
                raw_price = 0.0
                if price_col != -1 and price_col < len(r) and r[price_col] is not None:
                    try:
                        p_str = re.sub(r'[^0-9.]', '', str(r[price_col]))
                        if p_str:
                            raw_price = float(p_str)
                    except ValueError:
                        raw_price = 0.0

                if raw_price == 0.0 and invested_col != -1 and invested_col < len(r) and r[invested_col] is not None:
                    try:
                        inv_str = re.sub(r'[^0-9.]', '', str(r[invested_col]))
                        if inv_str and raw_qty > 0:
                            raw_price = round(float(inv_str) / raw_qty, 4)
                    except ValueError:
                        pass

                # 4. Parse Current Price / Current NAV
                current_price = raw_price
                if cur_price_col != -1 and cur_price_col < len(r) and r[cur_price_col] is not None:
                    try:
                        cp_str = re.sub(r'[^0-9.]', '', str(r[cur_price_col]))
                        if cp_str:
                            current_price = float(cp_str)
                    except ValueError:
                        current_price = raw_price

                if (current_price == 0.0 or current_price == raw_price) and val_col != -1 and val_col < len(r) and r[val_col] is not None:
                    try:
                        val_str = re.sub(r'[^0-9.]', '', str(r[val_col]))
                        if val_str and raw_qty > 0:
                            current_price = round(float(val_str) / raw_qty, 4)
                    except ValueError:
                        pass

                if current_price == 0.0 and raw_price > 0.0:
                    current_price = raw_price

                # 5. Parse Name
                asset_name = raw_sym_val
                if name_col != -1 and name_col < len(r) and r[name_col] is not None:
                    n_val = str(r[name_col]).strip()
                    if n_val:
                        asset_name = n_val

                self.upsert_asset(
                    symbol=clean_sym,
                    asset_type="MUTUAL_FUND" if "FUND" in clean_sym or "DIRECT" in clean_sym else "STOCK",
                    name=asset_name,
                    quantity=raw_qty,
                    buy_price=raw_price,
                    alert_threshold_pct=5.0,
                    notes="Imported from Demat / MF Statement (.xlsx)"
                )

                if current_price > 0:
                    conn = get_operational_db()
                    conn.execute(
                        "UPDATE user_portfolios SET current_price = ?, updated_at = ? WHERE asset_symbol = ?",
                        (current_price, datetime.now(timezone.utc).isoformat(), clean_sym)
                    )
                    conn.commit()
                    conn.close()

                imported_count += 1

            return imported_count
        except Exception as e:
            logger.error(f"Error parsing Excel Demat statement: {e}")
            return 0

    async def close(self):
        await self.http_client.aclose()
